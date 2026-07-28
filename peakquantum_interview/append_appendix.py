#!/usr/bin/env python3
"""Append two two-qubit-coupling concept slides to the existing deck's appendix."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from PIL import Image

ROOT = "/Users/nikolaygusarov/plasmon/peakquantum_interview"
FIN = f"{ROOT}/assets/final"
DECK = f"{ROOT}/Gusarov_PeakQuantum_QubitDesign.pptx"

BG=RGBColor(0x0C,0x14,0x22); CARD=RGBColor(0x16,0x22,0x39); LINE=RGBColor(0x2C,0x3E,0x5C)
COPPER=RGBColor(0xE0,0x8A,0x4C); COPPERL=RGBColor(0xF4,0xB0,0x72); TEAL=RGBColor(0x54,0xD8,0xC4)
TEXT=RGBColor(0xEE,0xF3,0xF8); MUTED=RGBColor(0x9F,0xB4,0xCC); PILL=RGBColor(0x20,0x2F,0x4A)
FONT="Calibri"

prs=Presentation(DECK)
BLANK=prs.slide_layouts[6]

def shadow(shape,blur=6,dist=3,dir_deg=90,alpha=0.42):
    xml=(f'<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
         f'<a:outerShdw blurRad="{int(blur*12700)}" dist="{int(dist*12700)}" dir="{int(dir_deg*60000)}" rotWithShape="0">'
         f'<a:srgbClr val="000000"><a:alpha val="{int(alpha*100000)}"/></a:srgbClr></a:outerShdw></a:effectLst>')
    shape._element.spPr.append(parse_xml(xml))

def panel(s,x,y,w,h,fill=CARD,radius=0.04,line=LINE,lw=0.75,sh=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.adjustments[0]=radius; sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False
    if sh: shadow(sp)
    return sp

def text(s,runs,x,y,w,h,size=14,color=TEXT,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=1.0):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,{})]
    p=tf.paragraphs[0]; p.alignment=align
    if spacing!=1.0: p.line_spacing=spacing
    for txt,o in runs:
        r=p.add_run(); r.text=txt; r.font.name=FONT; r.font.size=Pt(o.get("size",size))
        r.font.bold=o.get("bold",bold); r.font.color.rgb=o.get("color",color)
    return tb

def fit(s,path,bx,by,bw,bh):
    iw,ih=Image.open(path).size; ar=iw/ih
    if ar>bw/bh: w=bw; h=bw/ar
    else: h=bh; w=bh*ar
    return s.shapes.add_picture(path,Inches(bx+(bw-w)/2),Inches(by+(bh-h)/2),Inches(w),Inches(h))

def rrect(s,x,y,w,h,fill,radius=0.5,line=None,lw=1.0):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.adjustments[0]=radius; sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp

def legend(s,x,y,txt,sqcolor):
    sq=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y+0.05),Inches(0.16),Inches(0.16))
    sq.fill.solid(); sq.fill.fore_color.rgb=sqcolor; sq.line.fill.background(); sq.shadow.inherit=False
    text(s,[(txt,{"size":11,"color":MUTED})],x+0.24,y,2.6,0.28,anchor=MSO_ANCHOR.MIDDLE)

def bg(s):
    s.shapes.add_picture(f"{FIN}/bg.png",0,0,Inches(13.333),Inches(7.5))

def design_slide(title, sub, img, caption_runs, extra_note=None):
    s=prs.slides.add_slide(BLANK); bg(s)
    text(s,[(title,{"size":25,"bold":True,"color":TEXT})],0.6,0.36,12.1,0.55)
    text(s,[(sub,{"size":13,"color":COPPERL})],0.62,0.95,12.1,0.35)
    # legend top-right
    legend(s,9.9,0.5,"Nb — optical litho",COPPER)
    legend(s,11.75,0.5,"Al JJ — e-beam",TEAL)
    # framed drawing
    px,py,pw,ph=0.7,1.5,11.93,4.55
    panel(s,px,py,pw,ph,fill=RGBColor(0x11,0x1C,0x30),radius=0.03,line=LINE,lw=0.75)
    fit(s,img,px+0.25,py+0.2,pw-0.5,ph-0.4)
    # caption
    text(s,caption_runs,0.9,6.28,11.5,0.6,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,spacing=1.05)
    if extra_note:
        rrect(s,3.16,6.92,7.0,0.46,PILL,radius=0.5,line=LINE)
        text(s,[(extra_note,{"size":10.5,"color":TEAL})],3.3,6.92,6.72,0.46,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    return s

# Slide A — Plasmonium two-qubit
design_slide(
    "Two-qubit coupling — Plasmonium  (concept design)",
    "Answering: how would you couple two plasmonium qubits?",
    f"{FIN}/twoqubit_plasmonium.png",
    [("Two mergemon qubits, each capacitively coupled (C", {"size":12.5,"color":TEXT}),
     ("ᴄ", {"size":12.5,"color":TEXT}),
     (") to a shared ", {"size":12.5,"color":TEXT}),
     ("flux-tunable coupler SQUID", {"size":12.5,"bold":True,"color":COPPERL}),
     (" — coupling switched on/off for gates; each qubit keeps its own high-Z readout chain.",
      {"size":12.5,"color":TEXT})],
)

# Slide B — Fluxonium two-qubit
design_slide(
    "Two-qubit coupling — Fluxonium  (concept design)",
    "Current direction on the fluxonium project",
    f"{FIN}/twoqubit_fluxonium.png",
    [("Two fluxoniums (small JJ + JJ-array superinductor, capacitor-shunted) share a ", {"size":12.5,"color":TEXT}),
     ("coupling capacitor Cᴄ", {"size":12.5,"bold":True,"color":COPPERL}),
     (" — a fixed transverse coupling, the basis for a CZ / microwave two-qubit gate.", {"size":12.5,"color":TEXT})],
    extra_note="My role on fluxonium: HFSS EM design of a later chip iteration",
)

prs.save(DECK)
print("appended 2 slides ->", DECK, "| total slides:", len(prs.slides._sldIdLst))
